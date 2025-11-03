#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智慧水利平台架构与开发 - 课堂在线知识点导入转换器（修复版）
修复第5章和第6章解析问题
"""

import re
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def parse_markdown_tables_fixed(content):
    """修复版解析markdown内容中的所有表格"""
    lines = content.strip().split('\n')
    all_data_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # 找到表格开始
        if line.startswith('| 一级知识点'):
            # 跳过表头和分隔符行
            i += 2
            
            # 解析表格数据直到表格结束
            while i < len(lines):
                line = lines[i].strip()
                
                # 检查是否是有效的表格数据行
                if line.startswith('|') and line.endswith('|'):
                    # 排除分隔符行和表头重复行
                    if ('---' not in line and 
                        '一级知识点' not in line and 
                        '二级知识点' not in line):
                        
                        # 解析表格行
                        cells = [cell.strip() for cell in line.split('|')[1:-1]]
                        
                        if len(cells) >= 15:
                            # 更宽松的有效性检查
                            # 检查是否有任何实质内容（前7列有知识点，或后8列有属性）
                            has_knowledge = any(cell.strip() and cell.strip() != '' for cell in cells[:7])
                            has_attributes = any(cell.strip() and cell.strip() != '' for cell in cells[7:15])
                            
                            if has_knowledge or has_attributes:
                                all_data_rows.append(cells)
                
                # 检查表格结束条件
                elif (line.startswith('---') or 
                      line.startswith('>') or 
                      (line.startswith('##') and any(keyword in line for keyword in ['章', '节', '.'])) or 
                      line == ''):
                    # 遇到这些标志表格结束，跳出内循环继续寻找下一个表格
                    break
                
                i += 1
            continue
        i += 1
    
    return all_data_rows

def process_single_row_to_classroom_format(row):
    """将单行数据转换为课堂在线格式的多行数据"""
    if len(row) < 15:
        return []
    
    level1 = row[0].strip()
    level2 = row[1].strip()
    level3 = row[2].strip()
    level4 = row[3].strip()
    level5 = row[4].strip()
    level6 = row[5].strip()
    level7 = row[6].strip()
    
    pre_knowledge = row[7].strip()
    post_knowledge = row[8].strip()
    related_knowledge = row[9].strip()
    tags = row[10].strip()
    cognitive = row[11].strip()
    classification = row[12].strip()
    teaching_goal = row[13].strip()
    description = row[14].strip()
    
    result_rows = []
    
    # 分解多级知识点为单行，严格按照课堂在线要求
    if level1:
        # 一级知识点行
        new_row = [''] * 15
        new_row[0] = level1
        result_rows.append(new_row[:])
    
    if level2:
        # 二级知识点行
        new_row = [''] * 15
        new_row[1] = level2
        result_rows.append(new_row[:])
    
    if level3:
        # 三级知识点行
        new_row = [''] * 15
        new_row[2] = level3
        # 如果三级知识点有属性信息，且没有四级知识点，则在三级填充属性
        if (tags or cognitive or classification or teaching_goal or description) and not level4:
            new_row[7] = pre_knowledge
            new_row[8] = post_knowledge
            new_row[9] = related_knowledge
            new_row[10] = tags
            new_row[11] = cognitive
            new_row[12] = classification
            new_row[13] = teaching_goal
            new_row[14] = description
        result_rows.append(new_row[:])
    
    if level4:
        # 四级知识点行
        new_row = [''] * 15
        new_row[3] = level4
        # 四级知识点填充属性信息
        new_row[7] = pre_knowledge
        new_row[8] = post_knowledge
        new_row[9] = related_knowledge
        new_row[10] = tags
        new_row[11] = cognitive
        new_row[12] = classification
        new_row[13] = teaching_goal
        new_row[14] = description
        result_rows.append(new_row[:])
    
    # 处理五级知识点（如果存在）
    if level5 and level5.strip():
        new_row = [''] * 15
        new_row[4] = level5
        # 如果没有四级，五级可以填属性
        if not level4 and (tags or cognitive or classification or teaching_goal or description):
            new_row[7] = pre_knowledge
            new_row[8] = post_knowledge
            new_row[9] = related_knowledge
            new_row[10] = tags
            new_row[11] = cognitive
            new_row[12] = classification
            new_row[13] = teaching_goal
            new_row[14] = description
        result_rows.append(new_row[:])
    
    return result_rows

def convert_to_classroom_online_format(all_data_rows):
    """转换所有数据为课堂在线格式"""
    result_rows = []
    
    for row in all_data_rows:
        # 处理每一行，可能分解为多行
        converted_rows = process_single_row_to_classroom_format(row)
        result_rows.extend(converted_rows)
    
    return result_rows

def process_chapter_file_debug(file_path):
    """处理单个章节文件，带调试信息"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        file_name = file_path.name
        
        # 解析表格
        all_data_rows = parse_markdown_tables_fixed(content)
        
        print(f"  原始数据行数: {len(all_data_rows)}")
        
        # 显示前几行数据用于调试
        if len(all_data_rows) > 0:
            print(f"  前3行样例数据:")
            for i, row in enumerate(all_data_rows[:3]):
                knowledge_parts = [cell for cell in row[:7] if cell.strip()]
                attributes_parts = [cell for cell in row[7:15] if cell.strip()]
                print(f"    行{i+1}: 知识点={knowledge_parts}, 属性={len(attributes_parts)}个")
        
        # 转换为课堂在线格式
        converted_rows = convert_to_classroom_online_format(all_data_rows)
        
        return converted_rows
    
    except Exception as e:
        print(f"处理文件 {file_path} 时出错: {e}")
        import traceback
        traceback.print_exc()
        return []

def create_excel_file(all_rows, output_file):
    """创建Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "知识点导入"
    
    # 定义列名
    columns = [
        '一级知识点', '二级知识点', '三级知识点', '四级知识点', 
        '五级知识点', '六级知识点', '七级知识点',
        '前置知识点', '后置知识点', '关联知识点',
        '标签', '认知维度', '分类', '教学目标', '知识点说明'
    ]
    
    # 写入表头
    for col, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入数据
    for row_idx, row_data in enumerate(all_rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    # 设置列宽
    col_widths = {
        'A': 25, 'B': 25, 'C': 30, 'D': 35, 'E': 25, 'F': 20, 'G': 20,
        'H': 35, 'I': 35, 'J': 35, 'K': 25, 'L': 15, 'M': 15, 'N': 50, 'O': 70
    }
    
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 冻结表头
    ws.freeze_panes = 'A2'
    
    # 保存文件
    wb.save(output_file)

def main():
    """主函数"""
    base_dir = Path(__file__).parent
    
    # 定义章节文件
    chapter_files = [
        "第_1_章按小节的知识点清单（可直接粘贴到模板）.md",
        "第_2_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_3_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_4_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_5_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_6_章_超细化大纲式知识点表（严格空白对齐）.md"
    ]
    
    all_rows = []
    
    # 处理每个章节
    for chapter_file in chapter_files:
        file_path = base_dir / chapter_file
        if file_path.exists():
            print(f"\n处理文件: {chapter_file}")
            chapter_rows = process_chapter_file_debug(file_path)
            all_rows.extend(chapter_rows)
            print(f"  - 最终转换了 {len(chapter_rows)} 行数据")
        else:
            print(f"文件不存在: {chapter_file}")
    
    if not all_rows:
        print("没有找到有效的数据行")
        return
    
    # 输出Excel文件
    output_file = base_dir / "智慧水利平台架构与开发_课堂在线导入表格_修复版.xlsx"
    
    try:
        create_excel_file(all_rows, output_file)
        print(f"\n✅ 转换完成！")
        print(f"📊 总共处理了 {len(all_rows)} 行数据")
        print(f"📁 输出文件: {output_file}")
        
        # 统计各级知识点数量
        level_counts = [0] * 7
        for row in all_rows:
            for i in range(7):
                if row[i]:
                    level_counts[i] += 1
                    break
        
        print(f"\n📈 知识点分布统计：")
        for i, count in enumerate(level_counts):
            if count > 0:
                print(f"  {i+1}级知识点: {count} 个")
        
    except Exception as e:
        print(f"❌ 创建Excel文件时出错: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()