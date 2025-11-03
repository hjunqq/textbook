#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智慧水利平台架构与开发 - 课堂在线知识点导入转换器
将1-6章的markdown文件转换为符合课堂在线要求的Excel格式
"""

import re
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def parse_markdown_table(content):
    """解析markdown表格内容"""
    lines = content.strip().split('\n')
    data_rows = []
    
    # 分段解析，每个表格对应一个小节
    current_section = ""
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # 识别章节标题
        if line.startswith('##'):
            current_section = line.strip('# ').strip()
            i += 1
            continue
        
        # 找到表格开始
        if line.startswith('| 一级知识点'):
            # 跳过表头和分隔符行
            i += 2  # 跳过表头和分隔符
            
            # 解析表格数据
            while i < len(lines):
                line = lines[i].strip()
                if line.startswith('|') and line.endswith('|') and '---' not in line:
                    # 解析表格行
                    cells = [cell.strip() for cell in line.split('|')[1:-1]]
                    if len(cells) >= 15:
                        data_rows.append(cells)
                elif line.startswith('---') or line.startswith('>') or line.startswith('##') or not line.strip():
                    # 表格结束
                    break
                i += 1
            continue
        i += 1
    
    return data_rows

def convert_to_classroom_online_format(data_rows):
    """转换为课堂在线格式，确保每行只有一个知识点"""
    result_rows = []
    
    for row in data_rows:
        if len(row) < 15:
            continue
        
        # 提取各列数据
        level1 = row[0].strip()
        level2 = row[1].strip()
        level3 = row[2].strip()
        level4 = row[3].strip()
        level5 = row[4].strip()
        level6 = row[5].strip()
        level7 = row[6].strip()
        
        pre_knowledge = row[7].strip()
        post_knowledge = row[8].strip()
        related_knowledge = row[9].strip()
        tags = row[10].strip()
        cognitive = row[11].strip()
        classification = row[12].strip()
        teaching_goal = row[13].strip()
        description = row[14].strip()
        
        # 课堂在线格式要求：每行只能有一个知识点
        new_row = [''] * 15
        has_content = False
        
        # 按照层级关系，每行只填写一个层级的知识点
        if level1 and not level2 and not level3 and not level4:
            # 一级知识点单独成行
            new_row[0] = level1
            has_content = True
        elif level2 and not level3 and not level4:
            # 二级知识点单独成行
            new_row[1] = level2
            has_content = True
        elif level3 and not level4:
            # 三级知识点行
            new_row[2] = level3
            has_content = True
            # 三级知识点可以填写属性
            if any([tags, cognitive, classification, teaching_goal, description]):
                new_row[7] = pre_knowledge
                new_row[8] = post_knowledge
                new_row[9] = related_knowledge
                new_row[10] = tags
                new_row[11] = cognitive
                new_row[12] = classification
                new_row[13] = teaching_goal
                new_row[14] = description
        elif level4:
            # 四级知识点行
            new_row[3] = level4
            has_content = True
            # 四级知识点填写属性
            new_row[7] = pre_knowledge
            new_row[8] = post_knowledge
            new_row[9] = related_knowledge
            new_row[10] = tags
            new_row[11] = cognitive
            new_row[12] = classification
            new_row[13] = teaching_goal
            new_row[14] = description
        
        # 如果有有效内容，添加到结果中
        if has_content:
            result_rows.append(new_row)
    
    return result_rows

def process_chapter_file(file_path):
    """处理单个章节文件"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # 解析表格
        data_rows = parse_markdown_table(content)
        
        # 转换格式
        converted_rows = convert_to_classroom_online_format(data_rows)
        
        return converted_rows
    
    except Exception as e:
        print(f"处理文件 {file_path} 时出错: {e}")
        return []

def create_excel_file(all_rows, output_file):
    """创建Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "知识点导入"
    
    # 定义列名
    columns = [
        '一级知识点', '二级知识点', '三级知识点', '四级知识点', 
        '五级知识点', '六级知识点', '七级知识点',
        '前置知识点', '后置知识点', '关联知识点',
        '标签', '认知维度', '分类', '教学目标', '知识点说明'
    ]
    
    # 写入表头
    for col, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入数据
    for row_idx, row_data in enumerate(all_rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    # 设置列宽
    col_widths = {
        'A': 25, 'B': 25, 'C': 25, 'D': 25, 'E': 20, 'F': 20, 'G': 20,
        'H': 30, 'I': 30, 'J': 30, 'K': 20, 'L': 15, 'M': 15, 'N': 40, 'O': 60
    }
    
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 冻结表头
    ws.freeze_panes = 'A2'
    
    # 保存文件
    wb.save(output_file)

def main():
    """主函数"""
    base_dir = Path(__file__).parent
    
    # 定义章节文件
    chapter_files = [
        "第_1_章按小节的知识点清单（可直接粘贴到模板）.md",
        "第_2_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_3_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_4_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_5_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_6_章_超细化大纲式知识点表（严格空白对齐）.md"
    ]
    
    all_rows = []
    
    # 处理每个章节
    for chapter_file in chapter_files:
        file_path = base_dir / chapter_file
        if file_path.exists():
            print(f"处理文件: {chapter_file}")
            chapter_rows = process_chapter_file(file_path)
            all_rows.extend(chapter_rows)
            print(f"  - 转换了 {len(chapter_rows)} 行数据")
        else:
            print(f"文件不存在: {chapter_file}")
    
    if not all_rows:
        print("没有找到有效的数据行")
        return
    
    # 输出Excel文件
    output_file = base_dir / "智慧水利平台架构与开发_课堂在线导入表格.xlsx"
    
    try:
        create_excel_file(all_rows, output_file)
        print(f"\n转换完成！")
        print(f"总共处理了 {len(all_rows)} 行数据")
        print(f"输出文件: {output_file}")
        print(f"\n请检查生成的Excel文件，确保格式符合课堂在线的导入要求：")
        print("1. 每行只填写一个层级的知识点（A-G列中每行只有一列有内容）")
        print("2. 前置/后置/关联知识点用英文分号分隔")
        print("3. 标签用英文分号分隔")
        print("4. 认知维度和分类每个知识点只能填一个")
        print("5. 表头已冻结，方便查看")
    except Exception as e:
        print(f"创建Excel文件时出错: {e}")

if __name__ == "__main__":
    main()