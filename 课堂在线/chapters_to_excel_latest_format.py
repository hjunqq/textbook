#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智慧水利平台架构与开发 - 课堂在线知识点导入转换器（最新文档格式版）
解析 markdown 层级结构 + 属性块，适配2024新版课堂在线模板
"""

import re
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def parse_attributes(text):
    """解析属性块，支持 {key=value} 和 [key:value] 两种格式"""
    attributes = {
        'tags': '', 'cognitive': '', 'classification': '', 
        'goal': '', 'description': '', 'prerequisite': '', 
        'postrequisite': '', 'related': ''
    }
    
    # 提取属性块内容
    attr_match = re.search(r'[{\[]([^}\]]+)[}\]]', text)
    if not attr_match:
        return attributes, text
    
    attr_content = attr_match.group(1)
    clean_text = text[:attr_match.start()].strip()
    
    # 解析属性项（支持两种格式）
    # 格式1: key=value; 格式2: key:value
    attr_pairs = re.findall(r'([^;:=]+)[:=]([^;]+)', attr_content)
    
    for key, value in attr_pairs:
        key = key.strip().lower()
        value = value.strip()
        
        # 映射中文属性名
        key_mapping = {
            '标签': 'tags', 'tag': 'tags',
            '认知': 'cognitive', '认知维度': 'cognitive', 
            '分类': 'classification', '知识点分类': 'classification',
            '目标': 'goal', '教学目标': 'goal',
            '说明': 'description', '节点说明': 'description', '知识点说明': 'description',
            '前置': 'prerequisite', '前置知识点': 'prerequisite',
            '后置': 'postrequisite', '后置知识点': 'postrequisite',
            '关联': 'related', '关联知识点': 'related'
        }
        
        mapped_key = key_mapping.get(key, key)
        if mapped_key in attributes:
            attributes[mapped_key] = value
    
    return attributes, clean_text

def parse_markdown_hierarchy(content):
    """解析markdown层级结构"""
    lines = content.strip().split('\n')
    items = []
    
    for line in lines:
        line = line.rstrip()
        if not line or line.startswith('>') or line.startswith('---'):
            continue
            
        # 判断层级
        level = 0
        content_text = line
        
        if line.startswith('#'):
            # # 标题级别
            level = min(len(line) - len(line.lstrip('#')), 7)
            content_text = line.lstrip('#').strip()
        elif line.startswith('-') or line.startswith('  -'):
            # - 列表级别，根据缩进判断
            leading_spaces = len(line) - len(line.lstrip())
            if line.lstrip().startswith('-'):
                level = 3 + leading_spaces // 2  # 从第3级开始
                content_text = line.lstrip('- ').strip()
        
        if level > 0 and content_text:
            # 解析属性
            attributes, clean_content = parse_attributes(content_text)
            
            if clean_content:  # 只处理有内容的行
                items.append({
                    'level': level,
                    'content': clean_content,
                    'attributes': attributes
                })
    
    return items

def convert_to_new_template_format_v2(items):
    """转换为新模板格式（基于层级结构）"""
    result_rows = []
    path_tracker = {}  # 防止重复
    
    for item in items:
        level = item['level']
        content = item['content']
        attrs = item['attributes']
        
        # 跳过太长的内容（可能是章节说明）
        if len(content) > 200:
            continue
            
        # 创建路径键
        path_key = f"L{level}:{content}"
        if path_key in path_tracker:
            continue
        path_tracker[path_key] = True
        
        # 创建新行
        new_row = [''] * 14  # 14列：节点类型 + 7个节点名称 + 前置后置关联 + 标签分类说明
        
        # 判断节点类型
        has_attributes = any(attrs.values())
        is_knowledge_point = has_attributes and (attrs['tags'] or attrs['classification'] or attrs['goal'])
        
        if is_knowledge_point:
            new_row[0] = "知识点"
        else:
            new_row[0] = "分类"
        
        # 设置节点名称到对应级别（B到H列，索引1到7）
        if level <= 7:
            new_row[level] = content
        
        # 如果是知识点，填充属性
        if is_knowledge_point:
            new_row[8] = attrs['prerequisite']    # 前置节点
            new_row[9] = attrs['postrequisite']   # 后置节点
            new_row[10] = attrs['related']        # 关联节点
            new_row[11] = attrs['tags']           # 标签
            new_row[12] = attrs['classification'] # 知识点分类
            
            # 合并目标和说明
            description_parts = []
            if attrs['goal']:
                description_parts.append(f"目标: {attrs['goal']}")
            if attrs['description']:
                description_parts.append(f"说明: {attrs['description']}")
            new_row[13] = "; ".join(description_parts)
        
        result_rows.append(new_row)
    
    return result_rows

def create_excel_file_v2(all_rows, output_file):
    """创建Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "知识点导入"
    
    # 定义列名（按新模板格式）
    columns = [
        '节点类型*', '节点名称', '节点名称', '节点名称', '节点名称', 
        '节点名称', '节点名称', '节点名称', '前置节点', '后置节点', 
        '关联节点', '标签', '知识点分类', '节点说明'
    ]
    
    # 写入表头
    for col, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入数据
    for row_idx, row_data in enumerate(all_rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            if col_idx <= len(row_data):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    # 设置列宽
    col_widths = {
        'A': 12,  # 节点类型
        'B': 25, 'C': 25, 'D': 30, 'E': 35, 'F': 25, 'G': 20, 'H': 20,  # 7个节点名称列
        'I': 35, 'J': 35, 'K': 35,  # 前置后置关联
        'L': 25, 'M': 15, 'N': 70   # 标签分类说明
    }
    
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 冻结表头
    ws.freeze_panes = 'A2'
    
    # 保存文件
    wb.save(output_file)

def main():
    """主函数"""
    base_dir = Path(__file__).parent
    file_path = base_dir / "全章大纲_可直接导入（markdown_层级，细化版）.md"
    
    print("🚀 最新文档格式转换开始...")
    print("📋 解析 markdown 层级结构 + 属性块")
    print("🎯 适配2024版课堂在线新模板")
    
    if not file_path.exists():
        print(f"❌ 文件不存在: {file_path}")
        return
    
    try:
        # 读取文件
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        print(f"📖 处理文件: {file_path.name}")
        
        # 解析层级结构
        items = parse_markdown_hierarchy(content)
        print(f"  📋 解析出 {len(items)} 个节点")
        
        # 统计各级节点
        level_stats = {}
        attr_stats = {'knowledge_points': 0, 'categories': 0}
        
        for item in items:
            level = item['level']
            level_stats[f"L{level}"] = level_stats.get(f"L{level}", 0) + 1
            
            has_attributes = any(item['attributes'].values())
            if has_attributes:
                attr_stats['knowledge_points'] += 1
            else:
                attr_stats['categories'] += 1
        
        print(f"  📊 层级分布: {', '.join([f'{k}:{v}' for k, v in level_stats.items()])}")
        print(f"  🎯 预估: {attr_stats['categories']}个分类 + {attr_stats['knowledge_points']}个知识点")
        
        # 转换格式
        converted_rows = convert_to_new_template_format_v2(items)
        
        # 统计转换结果
        categories = sum(1 for row in converted_rows if row[0] == "分类")
        knowledge_points = sum(1 for row in converted_rows if row[0] == "知识点")
        
        print(f"  ✅ 转换完成: {categories}个分类 + {knowledge_points}个知识点 = {len(converted_rows)}个节点")
        
        if not converted_rows:
            print("❌ 没有生成有效的节点")
            return
        
        # 输出Excel文件
        output_file = base_dir / "智慧水利平台架构与开发_课堂在线导入表格_最新格式版.xlsx"
        
        create_excel_file_v2(converted_rows, output_file)
        
        print(f"\n🎉 最新格式转换完成！")
        print(f"📊 总共生成了 {len(converted_rows)} 个节点")
        print(f"📁 输出文件: {output_file}")
        
        print(f"\n📈 节点类型统计：")
        print(f"  📂 分类节点: {categories} 个")
        print(f"  📌 知识点: {knowledge_points} 个")
        
        # 模板限制检查
        print(f"\n🔍 模板限制检查：")
        print(f"  📊 节点总数: {len(converted_rows)}/5000 {'✅' if len(converted_rows) <= 5000 else '⚠️ 超限'}")
        
        # 估算关系数
        estimated_relations = sum(1 for row in converted_rows if any(row[8:11]))
        print(f"  🔗 关系估算: {estimated_relations}/2000 {'✅' if estimated_relations <= 2000 else '⚠️ 超限'}")
        
        print(f"\n✅ 最新格式版已生成！")
        print(f"📋 完全基于最新的markdown层级文档")
        print(f"🎯 结构清晰，属性完整，符合新模板要求")
        
        # 显示前几行示例
        print(f"\n📝 生成示例（前5行）:")
        for i, row in enumerate(converted_rows[:5]):
            node_type = row[0]
            node_name = next((row[j] for j in range(1, 8) if row[j]), "未知")
            tags = row[11] if row[11] else "无标签"
            print(f"  {i+1}. [{node_type}] {node_name} ({tags})")
        
    except Exception as e:
        print(f"❌ 处理过程中出错: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()