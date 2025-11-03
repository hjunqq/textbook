#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智慧水利平台架构与开发 - 课堂在线知识点导入转换器（增强版）
专门处理多表格的复杂markdown文件
"""

import re
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def debug_parse_markdown_tables(content, file_name):
    """调试解析markdown内容中的所有表格"""
    lines = content.strip().split('\n')
    all_data_rows = []
    table_count = 0
    
    print(f"\n=== 调试解析 {file_name} ===")
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # 找到表格开始
        if line.startswith('| 一级知识点'):
            table_count += 1
            print(f"找到第{table_count}个表格，起始行: {i+1}")
            
            # 跳过表头和分隔符行
            i += 2
            row_count = 0
            
            # 解析表格数据直到表格结束
            while i < len(lines):
                line = lines[i].strip()
                if line.startswith('|') and line.endswith('|') and '---' not in line:
                    # 解析表格行
                    cells = [cell.strip() for cell in line.split('|')[1:-1]]
                    if len(cells) >= 15:
                        # 检查是否有有效内容（前7列至少有一个非空）
                        if any(cell.strip() for cell in cells[:7]):
                            all_data_rows.append(cells)
                            row_count += 1
                elif line.startswith('---') or line.startswith('>') or line.startswith('##') or not line.strip():
                    # 表格结束
                    break
                i += 1
            
            print(f"  - 解析了{row_count}行有效数据")
            continue
        i += 1
    
    print(f"总计找到{table_count}个表格，解析了{len(all_data_rows)}行数据")
    return all_data_rows

def enhanced_parse_markdown_tables(content):
    """增强版解析markdown内容中的所有表格"""
    lines = content.strip().split('\n')
    all_data_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        
        # 找到表格开始
        if line.startswith('| 一级知识点'):
            # 跳过表头和分隔符行
            i += 2
            
            # 解析表格数据直到表格结束
            while i < len(lines):
                line = lines[i].strip()
                
                # 有效的表格数据行
                if (line.startswith('|') and line.endswith('|') and 
                    '---' not in line and 
                    line != '| 一级知识点 | 二级知识点 | 三级知识点 | 四级知识点 | 五级知识点 | 六级知识点 | 七级知识点 | 前置知识点 | 后置知识点 | 关联知识点 | 标签 | 认知维度 | 分类 | 教学目标 | 知识点说明 |'):
                    
                    # 解析表格行
                    cells = [cell.strip() for cell in line.split('|')[1:-1]]
                    
                    if len(cells) >= 15:
                        # 检查是否有有效内容（前7列至少有一个非空，或者有属性信息）
                        has_knowledge_point = any(cell.strip() for cell in cells[:7])
                        has_attributes = any(cell.strip() for cell in cells[7:15])
                        
                        if has_knowledge_point or has_attributes:
                            all_data_rows.append(cells)
                
                # 检查表格结束条件
                elif (line.startswith('---') or 
                      line.startswith('>') or 
                      (line.startswith('##') and any(keyword in line for keyword in ['章', '节', '.'])) or 
                      line == ''):
                    # 表格结束，但不要立即跳出，可能还有更多表格
                    break
                
                i += 1
            continue
        i += 1
    
    return all_data_rows

def process_single_row_to_classroom_format(row):
    """将单行数据转换为课堂在线格式的多行数据"""
    if len(row) < 15:
        return []
    
    level1 = row[0].strip()
    level2 = row[1].strip()
    level3 = row[2].strip()
    level4 = row[3].strip()
    level5 = row[4].strip()
    level6 = row[5].strip()
    level7 = row[6].strip()
    
    pre_knowledge = row[7].strip()
    post_knowledge = row[8].strip()
    related_knowledge = row[9].strip()
    tags = row[10].strip()
    cognitive = row[11].strip()
    classification = row[12].strip()
    teaching_goal = row[13].strip()
    description = row[14].strip()
    
    result_rows = []
    
    # 分解多级知识点为单行
    if level1:
        # 一级知识点行
        new_row = [''] * 15
        new_row[0] = level1
        result_rows.append(new_row[:])
    
    if level2:
        # 二级知识点行
        new_row = [''] * 15
        new_row[1] = level2
        result_rows.append(new_row[:])
    
    if level3:
        # 三级知识点行
        new_row = [''] * 15
        new_row[2] = level3
        # 如果三级知识点有属性信息，且没有四级知识点，则在三级填充属性
        if (tags or cognitive or classification or teaching_goal or description) and not level4:
            new_row[7] = pre_knowledge
            new_row[8] = post_knowledge
            new_row[9] = related_knowledge
            new_row[10] = tags
            new_row[11] = cognitive
            new_row[12] = classification
            new_row[13] = teaching_goal
            new_row[14] = description
        result_rows.append(new_row[:])
    
    if level4:
        # 四级知识点行
        new_row = [''] * 15
        new_row[3] = level4
        # 四级知识点填充属性信息
        new_row[7] = pre_knowledge
        new_row[8] = post_knowledge
        new_row[9] = related_knowledge
        new_row[10] = tags
        new_row[11] = cognitive
        new_row[12] = classification
        new_row[13] = teaching_goal
        new_row[14] = description
        result_rows.append(new_row[:])
    
    # 处理五级知识点（如果存在）
    if level5 and level5.strip():
        new_row = [''] * 15
        new_row[4] = level5
        # 五级知识点也可以填充属性信息
        if not level4:  # 如果没有四级，五级可以填属性
            new_row[7] = pre_knowledge
            new_row[8] = post_knowledge
            new_row[9] = related_knowledge
            new_row[10] = tags
            new_row[11] = cognitive
            new_row[12] = classification
            new_row[13] = teaching_goal
            new_row[14] = description
        result_rows.append(new_row[:])
    
    return result_rows

def convert_to_classroom_online_format(all_data_rows):
    """转换所有数据为课堂在线格式"""
    result_rows = []
    
    for row in all_data_rows:
        # 处理每一行，可能分解为多行
        converted_rows = process_single_row_to_classroom_format(row)
        result_rows.extend(converted_rows)
    
    return result_rows

def process_chapter_file(file_path, debug=False):
    """处理单个章节文件"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        file_name = file_path.name
        
        # 选择解析方法
        if debug:
            all_data_rows = debug_parse_markdown_tables(content, file_name)
        else:
            all_data_rows = enhanced_parse_markdown_tables(content)
        
        # 转换为课堂在线格式
        converted_rows = convert_to_classroom_online_format(all_data_rows)
        
        return converted_rows
    
    except Exception as e:
        print(f"处理文件 {file_path} 时出错: {e}")
        import traceback
        traceback.print_exc()
        return []

def create_excel_file(all_rows, output_file):
    """创建Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "知识点导入"
    
    # 定义列名
    columns = [
        '一级知识点', '二级知识点', '三级知识点', '四级知识点', 
        '五级知识点', '六级知识点', '七级知识点',
        '前置知识点', '后置知识点', '关联知识点',
        '标签', '认知维度', '分类', '教学目标', '知识点说明'
    ]
    
    # 写入表头
    for col, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入数据
    for row_idx, row_data in enumerate(all_rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    # 设置列宽
    col_widths = {
        'A': 25, 'B': 25, 'C': 30, 'D': 30, 'E': 20, 'F': 20, 'G': 20,
        'H': 35, 'I': 35, 'J': 35, 'K': 25, 'L': 15, 'M': 15, 'N': 50, 'O': 70
    }
    
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 冻结表头
    ws.freeze_panes = 'A2'
    
    # 保存文件
    wb.save(output_file)

def main():
    """主函数"""
    base_dir = Path(__file__).parent
    
    # 定义章节文件
    chapter_files = [
        "第_1_章按小节的知识点清单（可直接粘贴到模板）.md",
        "第_2_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_3_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_4_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_5_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_6_章_超细化大纲式知识点表（严格空白对齐）.md"
    ]
    
    all_rows = []
    
    # 处理每个章节
    for chapter_file in chapter_files:
        file_path = base_dir / chapter_file
        if file_path.exists():
            print(f"处理文件: {chapter_file}")
            
            # 对第5章和第6章进行调试解析
            debug_mode = '第_5_章' in chapter_file or '第_6_章' in chapter_file
            chapter_rows = process_chapter_file(file_path, debug=debug_mode)
            
            all_rows.extend(chapter_rows)
            print(f"  - 转换了 {len(chapter_rows)} 行数据")
        else:
            print(f"文件不存在: {chapter_file}")
    
    if not all_rows:
        print("没有找到有效的数据行")
        return
    
    # 输出Excel文件
    output_file = base_dir / "智慧水利平台架构与开发_课堂在线导入表格_完整版.xlsx"
    
    try:
        create_excel_file(all_rows, output_file)
        print(f"\n✅ 转换完成！")
        print(f"📊 总共处理了 {len(all_rows)} 行数据")
        print(f"📁 输出文件: {output_file}")
        
        # 统计各级知识点数量
        level_counts = [0] * 7
        for row in all_rows:
            for i in range(7):
                if row[i]:
                    level_counts[i] += 1
                    break
        
        print(f"\n📈 知识点分布统计：")
        for i, count in enumerate(level_counts):
            if count > 0:
                print(f"  {i+1}级知识点: {count} 个")
        
        # 统计各章节数量
        print(f"\n📝 各章节数据分布：")
        chapter_stats = {}
        
        # 重新统计各章节
        for chapter_file in chapter_files:
            file_path = base_dir / chapter_file
            if file_path.exists():
                chapter_rows = process_chapter_file(file_path, debug=False)
                chapter_name = chapter_file.split('_')[1].split('章')[0] + '章'
                chapter_stats[chapter_name] = len(chapter_rows)
        
        for chapter, count in chapter_stats.items():
            print(f"  {chapter}: {count} 行")
        
    except Exception as e:
        print(f"❌ 创建Excel文件时出错: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()