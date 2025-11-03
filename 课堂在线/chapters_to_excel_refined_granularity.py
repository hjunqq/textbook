#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智慧水利平台架构与开发 - 课堂在线知识点导入转换器（细分优化版）
进一步细分第三级节点，提高知识点颗粒度
"""

import re
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def parse_attributes(text):
    """解析属性块，支持 {key=value} 和 [key:value] 两种格式"""
    attributes = {
        'tags': '', 'cognitive': '', 'classification': '', 
        'goal': '', 'description': '', 'prerequisite': '', 
        'postrequisite': '', 'related': ''
    }
    
    # 提取属性块内容
    attr_match = re.search(r'[{\[]([^}\]]+)[}\]]', text)
    if not attr_match:
        return attributes, text
    
    attr_content = attr_match.group(1)
    clean_text = text[:attr_match.start()].strip()
    
    # 解析属性项（支持两种格式）
    attr_pairs = re.findall(r'([^;:=]+)[:=]([^;]+)', attr_content)
    
    for key, value in attr_pairs:
        key = key.strip().lower()
        value = value.strip()
        
        # 映射中文属性名
        key_mapping = {
            '标签': 'tags', 'tag': 'tags',
            '认知': 'cognitive', '认知维度': 'cognitive', 
            '分类': 'classification', '知识点分类': 'classification',
            '目标': 'goal', '教学目标': 'goal',
            '说明': 'description', '节点说明': 'description', '知识点说明': 'description',
            '前置': 'prerequisite', '前置知识点': 'prerequisite',
            '后置': 'postrequisite', '后置知识点': 'postrequisite',
            '关联': 'related', '关联知识点': 'related'
        }
        
        mapped_key = key_mapping.get(key, key)
        if mapped_key in attributes:
            attributes[mapped_key] = value
    
    return attributes, clean_text

def split_complex_content(content, level):
    """拆分复杂的内容节点"""
    # 移除常见的描述性前缀
    content = re.sub(r'^[^：:]+[:：]\s*', '', content)
    
    # 定义分隔符优先级（按优先级排序）
    separators = [';', '；', '/', '、']
    
    # 寻找最合适的分隔符
    for sep in separators:
        if sep in content:
            parts = [part.strip() for part in content.split(sep)]
            # 过滤掉太短或太长的部分
            parts = [p for p in parts if 2 <= len(p) <= 50 and not p.isdigit()]
            
            if len(parts) > 1:
                return parts
    
    return [content]

def parse_markdown_hierarchy_refined(content):
    """解析markdown层级结构（优化版）"""
    lines = content.strip().split('\n')
    items = []
    
    for line in lines:
        line = line.rstrip()
        if not line or line.startswith('>') or line.startswith('---'):
            continue
            
        # 判断层级
        level = 0
        content_text = line
        
        if line.startswith('#'):
            # # 标题级别
            level = min(len(line) - len(line.lstrip('#')), 7)
            content_text = line.lstrip('#').strip()
        elif line.startswith('-') or line.startswith('  -'):
            # - 列表级别，根据缩进判断
            leading_spaces = len(line) - len(line.lstrip())
            if line.lstrip().startswith('-'):
                level = 3 + leading_spaces // 2  # 从第3级开始
                content_text = line.lstrip('- ').strip()
        
        if level > 0 and content_text:
            # 解析属性
            attributes, clean_content = parse_attributes(content_text)
            
            if clean_content:  # 只处理有内容的行
                # 如果是第三级节点且包含复杂内容，尝试拆分
                if level == 3:
                    split_parts = split_complex_content(clean_content, level)
                    
                    if len(split_parts) > 1:
                        # 创建父分类节点（没有属性）
                        parent_name = re.sub(r'[:：].*', '', clean_content)
                        if parent_name != clean_content:
                            items.append({
                                'level': level,
                                'content': parent_name,
                                'attributes': {'tags': '', 'cognitive': '', 'classification': '', 'goal': '', 'description': '', 'prerequisite': '', 'postrequisite': '', 'related': ''},
                                'is_parent': True
                            })
                        
                        # 创建子知识点
                        for part in split_parts:
                            if part.strip():
                                items.append({
                                    'level': level + 1,
                                    'content': part.strip(),
                                    'attributes': attributes.copy(),
                                    'is_parent': False
                                })
                    else:
                        # 单个节点
                        items.append({
                            'level': level,
                            'content': clean_content,
                            'attributes': attributes,
                            'is_parent': False
                        })
                else:
                    # 其他级别正常处理
                    items.append({
                        'level': level,
                        'content': clean_content,
                        'attributes': attributes,
                        'is_parent': False
                    })
    
    return items

def convert_to_new_template_format_refined(items):
    """转换为新模板格式（细分优化版）"""
    result_rows = []
    path_tracker = {}  # 防止重复
    
    for item in items:
        level = item['level']
        content = item['content']
        attrs = item['attributes']
        is_parent = item.get('is_parent', False)
        
        # 跳过太长的内容（可能是章节说明）
        if len(content) > 200:
            continue
            
        # 创建路径键
        path_key = f"L{level}:{content}"
        if path_key in path_tracker:
            continue
        path_tracker[path_key] = True
        
        # 创建新行
        new_row = [''] * 14  # 14列：节点类型 + 7个节点名称 + 前置后置关联 + 标签分类说明
        
        # 判断节点类型
        has_attributes = any(attrs.values())
        is_knowledge_point = has_attributes and (attrs['tags'] or attrs['classification'] or attrs['goal']) and not is_parent
        
        if is_knowledge_point:
            new_row[0] = "知识点"
        else:
            new_row[0] = "分类"
        
        # 设置节点名称到对应级别（B到H列，索引1到7）
        if level <= 7:
            new_row[level] = content
        
        # 如果是知识点，填充属性
        if is_knowledge_point:
            new_row[8] = attrs['prerequisite']    # 前置节点
            new_row[9] = attrs['postrequisite']   # 后置节点  
            new_row[10] = attrs['related']        # 关联节点
            new_row[11] = attrs['tags']           # 标签
            new_row[12] = attrs['classification'] # 知识点分类
            
            # 合并目标和说明
            description_parts = []
            if attrs['goal']:
                description_parts.append(f"目标: {attrs['goal']}")
            if attrs['description']:
                description_parts.append(f"说明: {attrs['description']}")
            new_row[13] = "; ".join(description_parts)
        
        result_rows.append(new_row)
    
    return result_rows

def create_excel_file_refined(all_rows, output_file):
    """创建Excel文件（细分优化版）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "知识点导入"
    
    # 定义列名（按新模板格式）
    columns = [
        '节点类型*', '节点名称', '节点名称', '节点名称', '节点名称', 
        '节点名称', '节点名称', '节点名称', '前置节点', '后置节点', 
        '关联节点', '标签', '知识点分类', '节点说明'
    ]
    
    # 写入表头
    for col, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入数据
    for row_idx, row_data in enumerate(all_rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            if col_idx <= len(row_data):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    # 设置列宽
    col_widths = {
        'A': 12,  # 节点类型
        'B': 25, 'C': 25, 'D': 30, 'E': 35, 'F': 25, 'G': 20, 'H': 20,  # 7个节点名称列
        'I': 35, 'J': 35, 'K': 35,  # 前置后置关联
        'L': 25, 'M': 15, 'N': 70   # 标签分类说明
    }
    
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 冻结表头
    ws.freeze_panes = 'A2'
    
    # 保存文件
    wb.save(output_file)

def main():
    """主函数"""
    base_dir = Path(__file__).parent
    file_path = base_dir / "全章大纲_可直接导入（markdown_层级，细化版）.md"
    
    print("🚀 细分优化版转换开始...")
    print("📋 解析 markdown 层级结构 + 细分第三级复杂节点")
    print("🎯 提高知识点颗粒度，适配2024版课堂在线新模板")
    
    if not file_path.exists():
        print(f"❌ 文件不存在: {file_path}")
        return
    
    try:
        # 读取文件
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        print(f"📖 处理文件: {file_path.name}")
        
        # 解析层级结构（细分优化）
        items = parse_markdown_hierarchy_refined(content)
        print(f"  📋 解析出 {len(items)} 个节点（细分后）")
        
        # 统计各级节点
        level_stats = {}
        attr_stats = {'knowledge_points': 0, 'categories': 0}
        
        for item in items:
            level = item['level']
            level_stats[f"L{level}"] = level_stats.get(f"L{level}", 0) + 1
            
            has_attributes = any(item['attributes'].values())
            if has_attributes and not item.get('is_parent', False):
                attr_stats['knowledge_points'] += 1
            else:
                attr_stats['categories'] += 1
        
        print(f"  📊 层级分布: {', '.join([f'{k}:{v}' for k, v in level_stats.items()])}")
        print(f"  🎯 预估: {attr_stats['categories']}个分类 + {attr_stats['knowledge_points']}个知识点")
        
        # 转换格式
        converted_rows = convert_to_new_template_format_refined(items)
        
        # 统计转换结果
        categories = sum(1 for row in converted_rows if row[0] == "分类")
        knowledge_points = sum(1 for row in converted_rows if row[0] == "知识点")
        
        print(f"  ✅ 转换完成: {categories}个分类 + {knowledge_points}个知识点 = {len(converted_rows)}个节点")
        
        if not converted_rows:
            print("❌ 没有生成有效的节点")
            return
        
        # 输出Excel文件
        output_file = base_dir / "智慧水利平台架构与开发_课堂在线导入表格_细分优化版.xlsx"
        
        create_excel_file_refined(converted_rows, output_file)
        
        print(f"\n🎉 细分优化转换完成！")
        print(f"📊 总共生成了 {len(converted_rows)} 个节点")
        print(f"📁 输出文件: {output_file}")
        
        print(f"\n📈 节点类型统计：")
        print(f"  📂 分类节点: {categories} 个")
        print(f"  📌 知识点: {knowledge_points} 个")
        
        # 对比原始版本
        original_count = 192
        improvement = len(converted_rows) - original_count
        print(f"  🚀 相比原版增加: {improvement} 个节点 ({improvement/original_count*100:.1f}% 提升)")
        
        # 模板限制检查
        print(f"\n🔍 模板限制检查：")
        print(f"  📊 节点总数: {len(converted_rows)}/5000 {'✅' if len(converted_rows) <= 5000 else '⚠️ 超限'}")
        
        # 估算关系数
        estimated_relations = sum(1 for row in converted_rows if any(row[8:11]))
        print(f"  🔗 关系估算: {estimated_relations}/2000 {'✅' if estimated_relations <= 2000 else '⚠️ 超限'}")
        
        print(f"\n✅ 细分优化版已生成！")
        print(f"📋 第三级复杂节点已细分为独立知识点")
        print(f"🎯 知识点颗粒度大幅提升，学习路径更清晰")
        
        # 显示细分效果示例
        print(f"\n📝 细分效果示例（前10行）:")
        for i, row in enumerate(converted_rows[:10]):
            node_type = row[0]
            level_info = []
            for j in range(1, 8):
                if row[j]:
                    level_info.append(f"L{j}:{row[j]}")
            level_display = " > ".join(level_info) if level_info else "未知层级"
            tags = f" ({row[11]})" if row[11] else ""
            print(f"  {i+1}. [{node_type}] {level_display}{tags}")
        
    except Exception as e:
        print(f"❌ 处理过程中出错: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()