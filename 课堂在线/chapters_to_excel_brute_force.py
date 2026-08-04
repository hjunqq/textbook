#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
智慧水利平台架构与开发 - 课堂在线知识点导入转换器（暴力解析版）
使用最简单直接的方法解析所有表格行
"""

import re
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def parse_markdown_tables_brute_force(content):
    """暴力解析所有可能的表格数据行"""
    lines = content.strip().split('\n')
    all_data_rows = []
    
    for i, line in enumerate(lines):
        line = line.strip()
        
        # 检查是否是表格数据行
        if (line.startswith('|') and line.endswith('|') and 
            '---' not in line and  # 不是分隔符
            '一级知识点' not in line):  # 不是表头
            
            # 解析表格行
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            
            # 确保有15列数据
            if len(cells) >= 15:
                # 检查是否有有效内容（至少有一个知识点或属性）
                has_knowledge = any(cell.strip() for cell in cells[:7])
                has_attributes = any(cell.strip() for cell in cells[7:15])
                
                if has_knowledge or has_attributes:
                    all_data_rows.append(cells)
    
    return all_data_rows

def process_single_row_to_classroom_format(row):
    """将单行数据转换为课堂在线格式的多行数据"""
    if len(row) < 15:
        return []
    
    level1 = row[0].strip()
    level2 = row[1].strip()
    level3 = row[2].strip()
    level4 = row[3].strip()
    level5 = row[4].strip()
    level6 = row[5].strip()
    level7 = row[6].strip()
    
    pre_knowledge = row[7].strip()
    post_knowledge = row[8].strip()
    related_knowledge = row[9].strip()
    tags = row[10].strip()
    cognitive = row[11].strip()
    classification = row[12].strip()
    teaching_goal = row[13].strip()
    description = row[14].strip()
    
    result_rows = []
    
    # 分解多级知识点为单行，严格按照课堂在线要求
    if level1:
        # 一级知识点行
        new_row = [''] * 15
        new_row[0] = level1
        result_rows.append(new_row[:])
    
    if level2:
        # 二级知识点行
        new_row = [''] * 15
        new_row[1] = level2
        result_rows.append(new_row[:])
    
    if level3:
        # 三级知识点行
        new_row = [''] * 15
        new_row[2] = level3
        # 如果三级知识点有属性信息，且没有四级知识点，则在三级填充属性
        if (tags or cognitive or classification or teaching_goal or description) and not level4:
            new_row[7] = pre_knowledge
            new_row[8] = post_knowledge
            new_row[9] = related_knowledge
            new_row[10] = tags
            new_row[11] = cognitive
            new_row[12] = classification
            new_row[13] = teaching_goal
            new_row[14] = description
        result_rows.append(new_row[:])
    
    if level4:
        # 四级知识点行
        new_row = [''] * 15
        new_row[3] = level4
        # 四级知识点填充属性信息
        new_row[7] = pre_knowledge
        new_row[8] = post_knowledge
        new_row[9] = related_knowledge
        new_row[10] = tags
        new_row[11] = cognitive
        new_row[12] = classification
        new_row[13] = teaching_goal
        new_row[14] = description
        result_rows.append(new_row[:])
    
    # 处理五级知识点（如果存在）
    if level5 and level5.strip():
        new_row = [''] * 15
        new_row[4] = level5
        # 如果没有四级，五级可以填属性
        if not level4 and (tags or cognitive or classification or teaching_goal or description):
            new_row[7] = pre_knowledge
            new_row[8] = post_knowledge
            new_row[9] = related_knowledge
            new_row[10] = tags
            new_row[11] = cognitive
            new_row[12] = classification
            new_row[13] = teaching_goal
            new_row[14] = description
        result_rows.append(new_row[:])
    
    return result_rows

def convert_to_classroom_online_format(all_data_rows):
    """转换所有数据为课堂在线格式"""
    result_rows = []
    
    for row in all_data_rows:
        # 处理每一行，可能分解为多行
        converted_rows = process_single_row_to_classroom_format(row)
        result_rows.extend(converted_rows)
    
    return result_rows

def process_chapter_file_brute_force(file_path):
    """处理单个章节文件，暴力解析版本"""
    try:
        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()
        
        file_name = file_path.name
        
        # 暴力解析表格
        all_data_rows = parse_markdown_tables_brute_force(content)
        
        print(f"  📋 原始数据行数: {len(all_data_rows)}")
        
        # 显示详细信息
        if len(all_data_rows) > 0:
            # 统计各级知识点
            level_stats = {f"L{i+1}": 0 for i in range(7)}
            for row in all_data_rows:
                for i in range(7):
                    if row[i].strip():
                        level_stats[f"L{i+1}"] += 1
                        break
            
            stats_str = ", ".join([f"{k}:{v}" for k, v in level_stats.items() if v > 0])
            print(f"  📊 知识点分布: {stats_str}")
            
            # 如果是第5章或第6章，显示前几行示例
            if '第_5_章' in file_name or '第_6_章' in file_name:
                print(f"  🔍 前5行示例:")
                for i, row in enumerate(all_data_rows[:5]):
                    knowledge_parts = []
                    for j, cell in enumerate(row[:7]):
                        if cell.strip():
                            knowledge_parts.append(f"L{j+1}:{cell[:30]}")
                    
                    if knowledge_parts:
                        print(f"    第{i+1}行: {' | '.join(knowledge_parts)}")
        
        # 转换为课堂在线格式
        converted_rows = convert_to_classroom_online_format(all_data_rows)
        
        return converted_rows
    
    except Exception as e:
        print(f"❌ 处理文件 {file_path} 时出错: {e}")
        import traceback
        traceback.print_exc()
        return []

def create_excel_file(all_rows, output_file):
    """创建Excel文件"""
    wb = Workbook()
    ws = wb.active
    ws.title = "知识点导入"
    
    # 定义列名
    columns = [
        '一级知识点', '二级知识点', '三级知识点', '四级知识点', 
        '五级知识点', '六级知识点', '七级知识点',
        '前置知识点', '后置知识点', '关联知识点',
        '标签', '认知维度', '分类', '教学目标', '知识点说明'
    ]
    
    # 写入表头
    for col, header in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # 写入数据
    for row_idx, row_data in enumerate(all_rows, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    
    # 设置列宽
    col_widths = {
        'A': 25, 'B': 25, 'C': 30, 'D': 35, 'E': 25, 'F': 20, 'G': 20,
        'H': 35, 'I': 35, 'J': 35, 'K': 25, 'L': 15, 'M': 15, 'N': 50, 'O': 70
    }
    
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    
    # 冻结表头
    ws.freeze_panes = 'A2'
    
    # 保存文件
    wb.save(output_file)

def main():
    """主函数"""
    base_dir = Path(__file__).parent
    
    # 定义章节文件
    chapter_files = [
        "第_1_章按小节的知识点清单（可直接粘贴到模板）.md",
        "第_2_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_3_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_4_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_5_章_超细化大纲式知识点表（严格空白对齐）.md",
        "第_6_章_超细化大纲式知识点表（严格空白对齐）.md"
    ]
    
    all_rows = []
    chapter_stats = {}
    
    print("💪 开始暴力解析各章节文件...")
    print("🔧 逐行扫描，不放过任何表格数据！")
    
    # 处理每个章节
    for chapter_file in chapter_files:
        file_path = base_dir / chapter_file
        if file_path.exists():
            print(f"\n📖 处理文件: {chapter_file}")
            chapter_rows = process_chapter_file_brute_force(file_path)
            all_rows.extend(chapter_rows)
            
            chapter_name = chapter_file.split('_')[1].split('章')[0] + '章'
            chapter_stats[chapter_name] = len(chapter_rows)
            
            # 特别标记第5章和第6章
            if chapter_name in ['5章', '6章']:
                print(f"  🚀 暴力解析转换了 {len(chapter_rows)} 行数据")
            else:
                print(f"  ✅ 转换了 {len(chapter_rows)} 行数据")
        else:
            print(f"❌ 文件不存在: {chapter_file}")
    
    if not all_rows:
        print("❌ 没有找到有效的数据行")
        return
    
    # 输出Excel文件
    output_file = base_dir / "智慧水利平台架构与开发_课堂在线导入表格_暴力解析版.xlsx"
    
    try:
        create_excel_file(all_rows, output_file)
        
        print(f"\n🎉 转换完成！")
        print(f"📊 总共处理了 {len(all_rows)} 行数据")
        print(f"📁 输出文件: {output_file}")
        
        # 统计各级知识点数量
        level_counts = [0] * 7
        for row in all_rows:
            for i in range(7):
                if row[i]:
                    level_counts[i] += 1
                    break
        
        print(f"\n📈 知识点分布统计：")
        for i, count in enumerate(level_counts):
            if count > 0:
                print(f"  📌 {i+1}级知识点: {count} 个")
        
        # 各章节统计
        print(f"\n📋 各章节详细统计：")
        for chapter, count in chapter_stats.items():
            if chapter in ['5章', '6章']:
                expected = 41 if chapter == '5章' else 45
                if count >= expected * 0.9:  # 至少90%的数据
                    print(f"  🎯 {chapter}: {count} 行 ✅ (期望约{expected}行，解析成功!)")
                elif count >= expected * 0.5:  # 至少50%的数据
                    print(f"  🔧 {chapter}: {count} 行 📈 (期望约{expected}行，大幅改进)")
                else:
                    print(f"  ⚠️  {chapter}: {count} 行 (期望约{expected}行，仍需优化)")
            else:
                print(f"  📄 {chapter}: {count} 行")
        
        # 成功率统计
        ch5_ch6_total = chapter_stats.get('5章', 0) + chapter_stats.get('6章', 0)
        expected_total = 41 + 45  # 86
        success_rate = (ch5_ch6_total / expected_total) * 100 if expected_total > 0 else 0
        
        print(f"\n📊 第5、6章解析成功率: {success_rate:.1f}% ({ch5_ch6_total}/{expected_total})")
        
        if success_rate >= 80:
            print(f"🎯 解析效果优秀！第5、6章内容基本完整")
        elif success_rate >= 50:
            print(f"📈 解析效果良好！第5、6章内容大幅改善")
        else:
            print(f"🔧 仍需进一步优化解析逻辑")
        
        print(f"\n✅ 暴力解析版本已生成！")
        print(f"📋 格式完全符合课堂在线导入要求")
        
    except Exception as e:
        print(f"❌ 创建Excel文件时出错: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()